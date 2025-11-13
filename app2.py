import io
import zipfile
import datetime

import pandas as pd
import streamlit as st
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.worksheet.page import PageMargins

st.set_page_config(page_title="いわと発注管理", layout="centered")

TITLE = "いわと発注管理"

LEFT_HEADER_FONT = Font(name="ＭＳ ゴシック", size=28, bold=True)
CENTER_HEADER_FONT = Font(name="ＭＳ ゴシック", size=26, bold=True)
RIGHT_HEADER_FONT = Font(name="ＭＳ ゴシック", size=24, bold=True)
BODY_FONT = Font(name="ＭＳ ゴシック", size=22)
THIN = Side(border_style="thin", color="000000")


def style_sheet(ws):
    """本文・罫線・行高・列幅などの書式設定"""

    # 本文フォント・罫線・行高
    for row in ws.iter_rows(min_row=6):
        for c in row:
            c.font = BODY_FONT

            # デフォルトの配置
            align_kwargs = dict(
                vertical="center",
                wrap_text=True,
            )
            # B列だけ「縮小してセルに合わせる」をON
            if c.column_letter == "B":
                align_kwargs["shrink_to_fit"] = True

            c.alignment = Alignment(**align_kwargs)
            c.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

    # 行の高さ
    for i in range(1, ws.max_row + 1):
        ws.row_dimensions[i].height = 30

    # 6行目ヘッダー：フォントサイズ18＆上下左右中央
    header_font = Font(name="ＭＳ ゴシック", size=18)
    for cell in ws[6]:
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

    # 列幅設定
    ws.column_dimensions["A"].width = 15.18
    ws.column_dimensions["B"].width = 60.09
    for col in ["C", "D", "E", "F", "G", "H", "I", "J", "K", "L"]:
        ws.column_dimensions[col].width = 15.18

    # A4横設定
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_margins = PageMargins(left=0.3, right=0.3, top=0.5, bottom=0.5)


def set_header(ws, supplier, facility_label):
    """ヘッダー部分の書式"""

    # 左：仕入先 御中（A3:B3）
    ws.merge_cells("A3:B3")
    ws["A3"] = f"{supplier}　御中"
    ws["A3"].font = LEFT_HEADER_FONT
    ws["A3"].alignment = Alignment(horizontal="left", vertical="bottom")

    # 中央：注文書（施設名） → B1セルのみ（結合しない）
    ws["B1"] = f"注文書（{facility_label}）"
    ws["B1"].font = CENTER_HEADER_FONT
    ws["B1"].alignment = Alignment(horizontal="center", vertical="center")

    # 右：(有) ハートミール → K2セルに表示（結合しない）
    ws["K2"] = "(有) ハートミール"
    ws["K2"].font = RIGHT_HEADER_FONT
    ws["K2"].alignment = Alignment(horizontal="right", vertical="center")

    # 見出し近辺の行高
    ws.row_dimensions[1].height = 40
    ws.row_dimensions[2].height = 35
    ws.row_dimensions[3].height = 35


def blank_duplicated_dates_in_col_a(ws):
    """
    使用日：同じ日付が続く行はA列だけ空欄にする
    データは row 7 から始まる想定（row 6 がヘッダー）
    """
    prev_value = None
    for row in range(7, ws.max_row + 1):
        cell = ws.cell(row=row, column=1)  # A列
        if cell.value == prev_value:
            cell.value = None
        else:
            prev_value = cell.value


def ensure_columns(df, cols):
    for c in cols:
        if c not in df.columns:
            df[c] = ""
    return df


def forward_fill_cols(df, cols):
    for c in cols:
        if c in df.columns:
            df[c] = df[c].ffill()
    return df


def read_excel_flexible(file, header_row):
    # header_row は 1始まり想定 → pandasは0始まり
    hdr = max(0, header_row - 1)
    df = pd.read_excel(file, header=hdr)
    df.columns = df.columns.astype(str).str.strip().str.replace("\n", "", regex=False)
    return df


def to_excel_bytes(df, startrow=0):
    bio = io.BytesIO()
    with pd.ExcelWriter(bio, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, startrow=startrow)
    bio.seek(0)
    return bio.getvalue()


def output_order_excels_zip(df, facility):
    # 列の標準名を想定に寄せる（多少の揺れ吸収）
    rename_map = {}
    for c in df.columns:
        cc = str(c)
        if "使用日" in cc:
            rename_map[c] = "使用日"
        if "仕入先" in cc:
            rename_map[c] = "仕入先"
        if "食品名" in cc:
            rename_map[c] = "食品名"
        if "単位" in cc and "ユニ" not in cc:
            rename_map[c] = "単位"
        if "入所者" in cc and "ユ" not in cc:
            rename_map[c] = "入所者"
        if "職員" in cc:
            rename_map[c] = "職員"
        if "ユ" in cc and "入所者" in cc:
            rename_map[c] = "ユーハウス入所者"
        if "備考" in cc:
            rename_map[c] = "備考欄"
        if "納品時間" in cc:
            rename_map[c] = "納品時間"
        if "検収" in cc:
            rename_map[c] = "検収者"

    df = df.rename(columns=rename_map)

    # 欠損補完と型
    df = forward_fill_cols(df, ["使用日", "仕入先", "食品名"])

    if facility == "いわと":
        for c in ["入所者", "職員"]:
            if c in df.columns:
                df[c] = pd.to_numeric(df[c], errors="coerce").fillna(0)
    else:
        # ユーハウス
        if "ユーハウス入所者" in df.columns:
            df["ユーハウス入所者"] = pd.to_numeric(
                df["ユーハウス入所者"], errors="coerce"
            ).fillna(0)

    if "仕入先" not in df.columns:
        raise ValueError("「仕入先」列が見つかりません。ヘッダー行の指定を見直してください。")

    suppliers = df["仕入先"].dropna().unique()

    # 固定の列順（施設別で数量列が変わる）
    if facility == "いわと":
        keep_cols = [
            "使用日",
            "食品名",
            "入所者",
            "単位",
            "職員",
            "鮮度",
            "品温",
            "異物",
            "包装",
            "期限",
            "備考欄",
            "納品時間",
            "検収者",
        ]
        group_by = ["使用日", "食品名", "単位"]
        agg = {"入所者": "sum", "職員": "sum"}
        facility_label = "介護老人福祉施設いわと"
    else:
        keep_cols = [
            "使用日",
            "食品名",
            "ユーハウス入所者",
            "単位",
            "鮮度",
            "品温",
            "異物",
            "包装",
            "期限",
            "備考欄",
            "納品時間",
            "検収者",
        ]
        group_by = ["使用日", "食品名", "単位"]
        agg = {"ユーハウス入所者": "sum"}
        facility_label = "ユーハウスいわと"

    # ZIPに詰める
    zip_bytes = io.BytesIO()
    with zipfile.ZipFile(zip_bytes, mode="w", compression=zipfile.ZIP_DEFLATED) as zf:
        for supplier in suppliers:
            sub = df[df["仕入先"] == supplier].copy()

            # グループ集計
            present_keys = [k for k in group_by if k in sub.columns]
            sub = sub.groupby(present_keys, as_index=False).agg(agg)

            # 空列補完 → 並び替え
            sub = ensure_columns(sub, keep_cols)
            sub = sub[keep_cols]

            # 一旦Excelへ（startrow=5 でヘッダー余白）
            wb = Workbook()
            ws = wb.active

            # 先にテーブルを書き込む
            bio = io.BytesIO()
            with pd.ExcelWriter(bio, engine="openpyxl") as writer:
                sub.to_excel(writer, index=False, startrow=5)
            bio.seek(0)
            tmp_wb = load_workbook(bio)
            tmp_ws = tmp_wb.active

            # tmp_wsの内容をwsへコピー
            for r in tmp_ws.iter_rows(values_only=True):
                ws.append(list(r))

            # スタイル＆ヘッダー
            style_sheet(ws)
            set_header(ws, str(supplier), facility_label)

            # 使用日：同じ日付が続く行はA列だけ空欄にする
            blank_duplicated_dates_in_col_a(ws)

            # 仕入先別ファイル名
            safe = str(supplier).replace("/", "_").replace("\\", "_")
            out_name = f"注文書_{safe}_{facility}.xlsx"
            out_bio = io.BytesIO()
            wb.save(out_bio)
            out_bio.seek(0)
            zf.writestr(out_name, out_bio.read())

    zip_bytes.seek(0)
    return zip_bytes.getvalue()


# ====== UI ======
st.title(TITLE)
st.caption(
    "Python / Excel を会社PCに入れずに、ブラウザだけで『検収簿の加工→仕入先別注文書』を作れます。"
)

with st.expander("STEP 1：検収簿の加工（空欄補完付き）", expanded=True):
    st.write("※ 必要に応じて使います。既に“加工済みファイル”がある場合は STEP 2 へ。")

    f1 = st.file_uploader("検収記録簿（原本）をアップロード（.xlsx）", type=["xlsx"], key="raw")
    header_row = st.number_input(
        "ヘッダー（見出し）の行番号（1始まり）", min_value=1, value=1, step=1
    )

    cols_to_ffill = st.multiselect(
        "下方向にコピー（空欄補完）する列名",
        options=["納品日", "使用日", "朝昼夕", "仕入先"],
        # デフォルトは元コードの挙動を維持
        default=["納品日", "使用日", "朝昼夕", "仕入先"],
    )

    if st.button("加工する ▶", use_container_width=True, disabled=(f1 is None)):
        try:
            df_raw = read_excel_flexible(f1, header_row)
            df_proc = forward_fill_cols(df_raw, cols_to_ffill)

            ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
            st.success("加工が完了しました。ダウンロードしてください。")
            st.download_button(
                label="加工済ファイルをダウンロード",
                data=to_excel_bytes(df_proc, startrow=0),
                file_name=f"検収簿_加工済_空欄補完済み_{ts}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True,
            )
        except Exception as e:
            st.error(f"加工中にエラー：{e}")

st.markdown("---")

with st.expander("STEP 2：仕入先別 注文書を作成（ZIP）", expanded=True):
    f2 = st.file_uploader("加工済ファイルをアップロード（.xlsx）", type=["xlsx"], key="proc")

    facility = st.radio("施設を選択", options=["いわと", "ユーハウス"], horizontal=True)

    st.caption(
        "出力仕様：A4横／MSゴシック22pt／行高30／細罫線／"
        "ヘッダー（左：御中／中央：施設名／右：(有)ハートミール）／検収者列あり"
    )

    if st.button(
        "注文書を作成 ▶", use_container_width=True, disabled=(f2 is None)
    ):
        try:
            df2 = pd.read_excel(f2, header=0)
            zip_data = output_order_excels_zip(df2, facility=facility)
            ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
            st.success("仕入先別の注文書をZIPで用意しました。ダウンロードしてください。")
            st.download_button(
                label="注文書ZIPをダウンロード",
                data=zip_data,
                file_name=f"注文書_{facility}_{ts}.zip",
                mime="application/zip",
                use_container_width=True,
            )
        except Exception as e:
            st.error(f"作成中にエラー：{e}")
